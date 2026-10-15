"""Create journal-style DOCX submission files from the frozen Markdown and CSV assets."""

from __future__ import annotations

import csv
import re
from pathlib import Path

import latex2mathml.converter
import mathml2omml
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from xiangji.data import RESEARCH_ROOT
from xiangji.complete_manuscript import build as build_complete_manuscript


SUBMISSION_DIR = RESEARCH_ROOT / "submission"
MAIN_MD = RESEARCH_ROOT / "manuscript/manuscript_complete.md"
SOURCE_MD = RESEARCH_ROOT / "manuscript/manuscript_v1.md"
SUPPLEMENT_MD = RESEARCH_ROOT / "manuscript/supplementary_v1.md"
MAIN_TABLE_DIR = RESEARCH_ROOT / "manuscript/main_tables"
INK = RGBColor(31, 31, 29)

MAIN_TABLE_CAPTIONS = {
    1: "Notation used in the mathematical formulation.",
    2: "Development ablation results used for model selection. Chemistry was excluded from the locked primary decoder because it did not improve hidden-component F1.",
    3: "Corpus composition by historical work. The Hong work was quarantined for post-freeze confirmation.",
    4: "Primary source-macro comparison at 50% observed context. Each metric is averaged first within formula, then within work and finally across works.",
    5: "Work-level hidden-component F1 at 50% observed context. Work groups, rather than repeated masks, are the independent evaluation blocks.",
    6: "Sensitivity to the observed fraction of each formula. All entries use formula-first, work-second and equal-work aggregation.",
    7: "Uncertainty diagnostics for inclusion, ingredient count and composition. Ratio coverage is conditional on correct material selection.",
    8: "Fixed confirmation case study for Yinxiang fa at seed 17. Blank predicted ratios indicate materials absent from the predicted set.",
}

# Display numbers follow first appearance in the paper; source CSV numbers retain
# their semantic names so downstream analyses do not need to change.
MAIN_TABLE_SOURCE = {1: 1, 2: 4, 3: 2, 4: 3, 5: 5, 6: 6, 7: 7, 8: 8}
MAIN_FIGURE_SOURCE = {1: 1, 2: 2, 3: 3, 4: 6, 5: 4, 6: 5}

NUMBERED_SECTIONS = {
    "Introduction": 1,
    "Related Work": 2,
    "Problem Formulation and Dataset": 3,
    "Xiangji-CGIR Method": 4,
    "Experimental Design": 5,
    "Results": 6,
    "Discussion": 7,
    "Limitations": 8,
    "Conclusion": 9,
}

# Long equations are rendered as multiple native OMML rows in Word.  The LaTeX
# source remains a single numbered equation, so the mathematical content and
# numbering stay identical across formats.
WORD_EQUATION_LINES = {
    "composition-closure": [
        r"\mathbf{w}_i\in\Delta^{M-1},\qquad w_{ij}=\frac{q_{ij}}{\sum_{k\in S_i}q_{ik}},",
        r"S_i=\{j:w_{ij}>0\}.",
    ],
    "mask-cardinality": [
        r"r(K_i,f)=\min\{K_i-1,\max[1,\operatorname{round}(fK_i)]\},",
        r"|O_i|=r(K_i,f).",
    ],
    "prevalence-priors": [
        r"\pi_j=\frac{c_j+1/2}{N+1},",
        r"\pi_{j\mid d}=\frac{c_{dj}+\lambda_s\pi_j}{N_d+\lambda_s},\qquad\lambda_s=5.",
    ],
    "graph-affinity": [
        r"O_i=\{o_{i1},\ldots,o_{i|O_i|}\},",
        r"a^{(G)}_{ij}=\frac{1}{|O_i|}\sum_{\ell=1}^{|O_i|}G_{o_{i\ell}j}.",
    ],
    "title-affinity": [
        r"\mathbf u_j=\frac{\sum_{r=1}^{N}y_{rj}\mathbf v(t_r)}{\|\sum_{r=1}^{N}y_{rj}\mathbf v(t_r)\|_2},",
        r"a^{(T)}_{ij}=\mathbf v(t_i)^{\top}\mathbf u_j.",
    ],
    "candidate-features": [
        r"g_{ij}=\operatorname{logit}\pi_{j\mid d_i}-\operatorname{logit}\pi_j,",
        r"u_{ij}={\bf1}(j\in O_i),",
        r"\boldsymbol\phi_{ij}=[\operatorname{logit}\pi_j,g_{ij},a^{(T)}_{ij},a^{(G)}_{ij}]^{\top},",
        r"\mathbf x_{ij}=[\boldsymbol\phi_{ij}^{\top},u_{ij},u_{ij}\log w_{ij}]^{\top}.",
    ],
    "logistic-model": [
        r"p_n(\boldsymbol\theta)=\sigma(\tilde{\mathbf z}_n^{\top}\boldsymbol\theta),",
        r"\sigma(u)=\frac{1}{1+\exp(-u)}.",
    ],
    "map-objective": [
        r"\ell_n(\boldsymbol\theta)=\omega_n[\log(1+\exp\eta_n)-y_n\eta_n],",
        r"\mathcal L_{\mathrm{data}}(\boldsymbol\theta)=\sum_{n=1}^{L_{\mathrm{tr}}}\ell_n(\boldsymbol\theta),",
        r"\mathcal L(\boldsymbol\theta)=\mathcal L_{\mathrm{data}}(\boldsymbol\theta)+\tfrac12\boldsymbol\theta^{\top}\boldsymbol\Lambda\boldsymbol\theta,",
        r"\widehat{\boldsymbol\theta}=\underset{\boldsymbol\theta}{\mathrm{arg\,min}}\,\mathcal L(\boldsymbol\theta),\qquad\eta_n=\tilde{\mathbf z}_n^{\top}\boldsymbol\theta.",
    ],
    "laplace-posterior": [
        r"\boldsymbol\Omega=\operatorname{diag}\{\omega_n\widehat p_n(1-\widehat p_n)\},",
        r"\boldsymbol\Sigma=[\tilde{\mathbf Z}^{\top}\boldsymbol\Omega\tilde{\mathbf Z}+\boldsymbol\Lambda]^+,",
        r"\boldsymbol\theta\mid\mathcal D\approx\mathcal N(\widehat{\boldsymbol\theta},\boldsymbol\Sigma).",
    ],
    "predictive-probability": [
        r"v_{ij}=\tilde{\mathbf z}_{ij}^{\top}\boldsymbol\Sigma\tilde{\mathbf z}_{ij},",
        r"s_{ij}=(1+\pi v_{ij}/8)^{1/2},",
        r"\widehat p_{ij}=\sigma[(\tilde{\mathbf z}_{ij}^{\top}\widehat{\boldsymbol\theta}+\delta)/s_{ij}].",
    ],
    "count-priors": [
        r"p_0(k)=\frac{n_k+1/2}{N+K_{\max}/2},",
        r"p_d(k)=\frac{n_{dk}+\lambda_kp_0(k)}{N_d+\lambda_k},\qquad\lambda_k=3.",
    ],
    "count-posterior": [
        r"\mathcal K_i(o,f)=\{k:k\ge o,\ r(k,f)=o\},",
        r"P_i(k)=p(K_i=k\mid d_i,o,f),",
        r"P_i(k)\propto p_{d_i}(k){\bf1}\{k\in\mathcal K_i(o,f)\}.",
    ],
    "count-decision": [
        r"F_i(k)=\sum_{h=1}^{k}p(K_i=h\mid d_i,o,f),",
        r"\widehat K_i=\min\{k:F_i(k)\ge\tfrac12\}.",
    ],
    "ratio-ridge": [
        r"\mathcal R(\mathbf b)=\sum_{n=1}^{L_r}(\ell_n-\tilde{\mathbf z}_n^{\top}\mathbf b)^2+\rho\|\mathbf b\|_2^2,",
        r"\widehat{\mathbf b}=\underset{\mathbf b}{\mathrm{arg\,min}}\,\mathcal R(\mathbf b).",
    ],
    "ratio-closure": [
        r"r_{ij}=\exp(\tilde{\mathbf z}_{ij}^{\top}\widehat{\mathbf b}),",
        r"\widehat S_i=\{s_{i1},\ldots,s_{i|\widehat S_i|}\},",
        r"\mu_{ij}=r_{ij}/\sum_{h=1}^{|\widehat S_i|}r_{i,s_{ih}},\qquad j\in\widehat S_i.",
    ],
    "mass-constraint": [
        r"H_i=\{h_{i1},\ldots,h_{i|H_i|}\},",
        r"h\in H_i\Leftrightarrow h\in\widehat S_i\ \mathrm{and}\ h\notin O_i,",
        r"R_i=1-\sum_{\ell=1}^{|O_i|}w_{i,o_{i\ell}},\qquad Z_i=\sum_{\ell=1}^{|H_i|}\mu_{i,h_{i\ell}},",
        r"\widehat w^{(c)}_{ij}=w_{ij}\quad(j\in O_i),",
        r"\widehat w^{(c)}_{ij}=R_i\mu_{ij}/Z_i\quad(j\in H_i),",
        r"\widehat w^{(c)}_{ij}=0\quad(\mathrm{otherwise}).",
    ],
    "profile-similarities": [
        r"T(\mathbf f_p,\mathbf f_q)=\frac{\sum_{b=1}^{B}\min(f_{pb},f_{qb})}{\sum_{b=1}^{B}\max(f_{pb},f_{qb})},",
        r"C_+(\mathbf a,\mathbf b)=\max\left(0,\frac{\mathbf a^{\top}\mathbf b}{\|\mathbf a\|_2\|\mathbf b\|_2}\right).",
    ],
    "chemical-kernel": [
        r"c^{(r)}_{pq}=C_+(\mathbf r_p,\mathbf r_q),",
        r"c^{(o)}_{pq}=C_+(\mathbf o_p,\mathbf o_q),",
        r"c^{(\Sigma)}_{pq}=c^{(r)}_{pq}+c^{(o)}_{pq},",
        r"\kappa(p,q)=0.50T(\mathbf f_p,\mathbf f_q)+0.25c^{(\Sigma)}_{pq},",
        r"K^{(C)}_{ab}=\max_{p\in\mathcal P_a,q\in\mathcal P_b}\kappa(p,q).",
    ],
    "knn-similarity": [
        r"s_{ir}=\frac{\mathbf o_i^{\top}\mathbf y_r}{\|\mathbf o_i\|_2\|\mathbf y_r\|_2},",
        r"\lambda_{ir}=s_{ir}/\sum_{h=1}^{10}s_{ih},\qquad\widehat{\mathbf p}_i=\sum_{r=1}^{10}\lambda_{ir}\mathbf y_r.",
    ],
    "set-metrics": [
        r"\operatorname{Prec}_i=\frac{|H_i\cap P_i|}{|P_i|},\qquad\operatorname{Rec}_i=\frac{|H_i\cap P_i|}{|H_i|},",
        r"F_{1i}=\frac{2\operatorname{Prec}_i\operatorname{Rec}_i}{\operatorname{Prec}_i+\operatorname{Rec}_i}.",
    ],
    "composition-metrics": [
        r"J_{wi}=\frac{\sum_{\ell=1}^{|U_i|}\min(w_{i\ell},\widehat w_{i\ell})}{\sum_{\ell=1}^{|U_i|}\max(w_{i\ell},\widehat w_{i\ell})},",
        r"E_{Ui}=\frac{1}{|U_i|}\sum_{\ell=1}^{|U_i|}|w_{i\ell}-\widehat w_{i\ell}|.",
    ],
    "aitchison-distance": [
        r"d_{Ai}=|U_i|^{-1/2}\|\operatorname{clr}(\mathbf w_i^{\epsilon})-\operatorname{clr}(\widehat{\mathbf w}_i^{\epsilon})\|_2,",
        r"\operatorname{clr}(\mathbf x)=\log\mathbf x-|U_i|^{-1}\sum_{\ell=1}^{|U_i|}\log x_{\ell}.",
    ],
    "calibration-metrics": [
        r"\operatorname{Brier}=\frac{1}{L}\sum_{n=1}^{L}(p_n-y_n)^2,",
        r"\operatorname{ECE}=\sum_{b=1}^{10}\frac{|B_b|}{L}|\overline p_b-\overline y_b|.",
    ],
    "source-macro": [
        r"\overline m_i=|A_i|^{-1}\sum_{a=1}^{|A_i|}m_{ia},",
        r"\overline m_d=|I_d|^{-1}\sum_{i=1}^{|I_d|}\overline m_i,",
        r"m_{\mathrm{macro}}=\frac{1}{D}\sum_{d=1}^{D}\overline m_d.",
    ],
    "sign-flip-test": [
        r"T(\mathbf s)=\left|D^{-1}\sum_{d=1}^{D}s_d\Delta_d\right|,",
        r"\mathcal S=\{\mathbf s_1,\ldots,\mathbf s_{2^D}\}=\{-1,1\}^{D},",
        r"p_{\mathrm{flip}}=2^{-D}\sum_{h=1}^{2^D}{\bf1}\{T(\mathbf s_h)\ge T(\mathbf1)\}.",
    ],
    "count-conformal": [
        r"q=\operatorname{Quantile}_{\lceil(n+1)0.8\rceil/n}^{\mathrm{higher}}(a_1,\ldots,a_n),",
        r"L_i^{(c)}=\max(|O_i|+1,L_i-q),",
        r"U_i^{(c)}=U_i+q.",
    ],
}


def _romanize_for_portable_pdf(text: str) -> str:
    replacements = (
        ("印香法 (*Yinxiang fa*)", "*Yinxiang fa*"),
        ("印香法", "Yinxiang fa"), ("陳氏香譜", "Chen shi xiangpu"),
        ("四庫全書本", "Siku Quanshu edition"), ("香乘", "Xiangcheng"),
        ("香譜", "Xiangpu"), ("卷下", "juan xia"), ("卷", "juan "),
    )
    for source, target in replacements:
        text = text.replace(source, target)
    if re.search(r"[\u3400-\u9fff]", text):
        raise ValueError(f"Unmapped CJK text in portable submission export: {text}")
    return text


def _set_cell_margins(cell, top=70, start=80, bottom=70, end=80) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def _set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    old = properties.find(qn("w:tblBorders"))
    if old is not None:
        properties.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge, value, size in (
        ("top", "single", "10"), ("bottom", "single", "10"),
        ("left", "nil", "0"), ("right", "nil", "0"),
        ("insideH", "nil", "0"), ("insideV", "nil", "0"),
    ):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), value)
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), "000000")
        borders.append(node)
    properties.append(borders)


def _remove_table_borders(table) -> None:
    properties = table._tbl.tblPr
    old = properties.find(qn("w:tblBorders"))
    if old is not None:
        properties.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    properties.append(borders)


def _set_header_bottom_border(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)


def _enable_line_numbers(section) -> None:
    properties = section._sectPr
    node = properties.find(qn("w:lnNumType"))
    if node is None:
        node = OxmlElement("w:lnNumType")
        properties.append(node)
    node.set(qn("w:countBy"), "1")
    node.set(qn("w:restart"), "continuous")


def _page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def _configure(document: Document, *, line_numbers: bool) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.8)
    if line_numbers:
        _enable_line_numbers(section)
    _page_number(section.footer.paragraphs[0])

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    for name, size, before, after in (("Title", 18, 0, 14), ("Heading 1", 13.5, 16, 7), ("Heading 2", 11.5, 12, 4)):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        style.font.color.rgb = INK
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        fonts = style._element.rPr.rFonts
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            key = qn(f"w:{attribute}")
            if key in fonts.attrib:
                del fonts.attrib[key]
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        if name == "Title":
            borders = style._element.pPr.find(qn("w:pBdr"))
            if borders is not None:
                style._element.pPr.remove(borders)

    if "Reference" not in [style.name for style in document.styles]:
        style = document.styles.add_style("Reference", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.size = Pt(9)
        style.paragraph_format.left_indent = Inches(0.25)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_after = Pt(3)

    if "Code" not in [style.name for style in document.styles]:
        style = document.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Menlo"
        style.font.size = Pt(8)
        style.paragraph_format.left_indent = Inches(0.25)
        style.paragraph_format.space_after = Pt(1)

    for name in ("Hyperlink", "FollowedHyperlink"):
        if name in [style.name for style in document.styles]:
            style = document.styles[name]
            style.font.name = "Times New Roman"
            style.font.color.rgb = INK
            style.font.underline = False

    # Scrub the complete style sheet, including unused built-in headings.  Some
    # Word templates retain theme-blue defaults in styles that never appear on
    # the page; keeping the package itself monochrome prevents later edits from
    # unexpectedly reintroducing blue text.
    for style in document.styles:
        properties = style._element.get_or_add_rPr()
        fonts = properties.get_or_add_rFonts()
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attribute}"), "Times New Roman")
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{attribute}"), None)
        color = properties.find(qn("w:color"))
        if color is None:
            color = OxmlElement("w:color")
            properties.append(color)
        color.set(qn("w:val"), "1F1F1D")
        for attribute in ("themeColor", "themeTint", "themeShade"):
            color.attrib.pop(qn(f"w:{attribute}"), None)
        for node in style._element.iter():
            for key, value in list(node.attrib.items()):
                if value.upper() == "4F81BD":
                    node.set(key, "1F1F1D")
            for attribute in ("themeColor", "themeTint", "themeShade"):
                node.attrib.pop(qn(f"w:{attribute}"), None)


def _omml(latex: str):
    # ``mathml2omml`` emits malformed grouping XML for ``\widetilde``.  Word's
    # native equation renderer treats ``\tilde`` identically for these vectors.
    latex = latex.replace(r"\widetilde", r"\tilde")
    mathml = latex2mathml.converter.convert(latex)
    converted = mathml2omml.convert(mathml)
    wrapper = parse_xml(
        '<root xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        + converted + "</root>"
    )
    equation = wrapper[0]
    math_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    for math_run in equation.findall(f".//{{{math_namespace}}}r"):
        word_properties = math_run.find(qn("w:rPr"))
        if word_properties is None:
            word_properties = OxmlElement("w:rPr")
            math_run.insert(1 if len(math_run) and math_run[0].tag == f"{{{math_namespace}}}rPr" else 0, word_properties)
        for tag in ("sz", "szCs"):
            node = word_properties.find(qn(f"w:{tag}"))
            if node is None:
                node = OxmlElement(f"w:{tag}")
                word_properties.append(node)
            node.set(qn("w:val"), "16")
    return equation


def _add_display_equation(document: Document, label: str, latex: str, number: str) -> None:
    lines = WORD_EQUATION_LINES.get(label, [latex])
    table = document.add_table(rows=1, cols=3)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _remove_table_borders(table)
    widths = (Inches(0.42), Inches(5.96), Inches(0.42))
    row = table.rows[0]
    for cell, width in zip(row.cells, widths):
        cell.width = width
        _set_cell_margins(cell, top=5, start=0, bottom=5, end=0)
    equation_cell = table.cell(0, 1)
    for line_index, expression in enumerate(lines):
        equation_paragraph = equation_cell.paragraphs[0] if line_index == 0 else equation_cell.add_paragraph()
        equation_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        equation_paragraph.paragraph_format.space_after = Pt(0)
        equation_paragraph.paragraph_format.keep_with_next = line_index < len(lines) - 1
        equation_paragraph._p.append(_omml(expression))
    number_cell = table.cell(0, 2)
    number_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    number_paragraph = number_cell.paragraphs[0]
    number_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    number_paragraph.paragraph_format.space_after = Pt(0)
    run = number_paragraph.add_run(f"({number})")
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _figure_captions() -> dict[int, str]:
    text = SOURCE_MD.read_text(encoding="utf-8")
    legend_text = text.split("## Figure legends", 1)[1]
    captions: dict[int, str] = {}
    pattern = re.compile(r"^\*\*Figure (\d+) \| (.+?)\.\*\*\s*(.*)$", re.MULTILINE)
    for match in pattern.finditer(legend_text):
        number = int(match.group(1))
        captions[number] = f"{match.group(2)}. {match.group(3)}".strip()
    if set(captions) != set(range(1, 7)):
        raise ValueError(f"Expected six figure captions, found {sorted(captions)}")
    return captions


def _add_main_figure(document: Document, number: int) -> None:
    source_number = MAIN_FIGURE_SOURCE[number]
    path = next((RESEARCH_ROOT / "figures").glob(f"fig{source_number}_*.png"))
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run().add_picture(str(path), width=Inches(6.35))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    caption.paragraph_format.space_after = Pt(9)
    label = caption.add_run(f"Figure {number}. ")
    label.bold = True
    _inline(caption, _figure_captions()[source_number])
    for run in caption.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)
        run.font.color.rgb = INK


def _add_main_table(document: Document, number: int) -> None:
    path = next(MAIN_TABLE_DIR.glob(f"table{MAIN_TABLE_SOURCE[number]}_*.csv"))
    with path.open(encoding="utf-8", newline="") as handle:
        rows = [[_romanize_for_portable_pdf(value) for value in row] for row in csv.reader(handle)]
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True
    label = caption.add_run(f"Table {number}. ")
    label.bold = True
    _inline(caption, MAIN_TABLE_CAPTIONS[number])
    for run in caption.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)
        run.font.color.rgb = INK

    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _set_table_borders(table)
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell, top=45, start=55, bottom=45, end=55)
            cell.text = value
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.keep_with_next = row_index < len(rows) - 1
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(7.7 if len(rows[0]) >= 4 else 8.5)
                    run.font.color.rgb = INK
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            if row_index == 0:
                _set_header_bottom_border(cell)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
        table.rows[row_index]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            table.rows[row_index]._tr.get_or_add_trPr().append(header)


def _inline(paragraph, text: str) -> None:
    token_pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\$[^$]+\$)")
    position = 0
    for match in token_pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Menlo"
            run.font.size = Pt(8.5)
        else:
            paragraph._p.append(_omml(token[1:-1]))
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.color.rgb = INK
        fonts = run._r.get_or_add_rPr().rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            run._r.get_or_add_rPr().insert(0, fonts)
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attribute}"), "Times New Roman")


def _add_markdown(
    document: Document, path: Path, *, stop_before_tables: bool = False,
    equation_prefix: str = "",
) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    in_code = False
    buffer: list[str] = []
    in_references = False
    equation_number = 0
    current_section_number: int | None = None
    subsection_number = 0

    def flush() -> None:
        nonlocal buffer
        if not buffer:
            return
        text = _romanize_for_portable_pdf(" ".join(value.strip() for value in buffer).strip())
        buffer = []
        if not text:
            return
        if "[Author name]" in text and "[Affiliation, city, country]" in text:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(12)
            for index, value in enumerate(("[Author name]", "[Affiliation, city, country]", "[Corresponding email]")):
                _inline(paragraph, value)
                if index < 2:
                    paragraph.add_run().add_break()
            return
        style = "Reference" if in_references and re.match(r"^\d+\.", text) else None
        paragraph = document.add_paragraph(style=style)
        _inline(paragraph, text)

    for raw in lines:
        line = raw.rstrip()
        if stop_before_tables and line == "### Supplementary Tables":
            flush()
            break
        if line.startswith("```"):
            flush()
            in_code = not in_code
            continue
        if in_code:
            paragraph = document.add_paragraph(style="Code")
            paragraph.add_run(line)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        equation = re.match(r"^\$\$\[([a-z0-9-]+)\]\s*(.+?)\s*\$\$$", line)
        marker = re.match(r"^\[\[(FIGURE|TABLE):(\d+)\]\]$", line)
        if marker:
            flush()
            kind, number = marker.group(1), int(marker.group(2))
            if kind == "FIGURE":
                _add_main_figure(document, number)
            else:
                _add_main_table(document, number)
            continue
        if equation:
            flush()
            equation_number += 1
            _add_display_equation(document, equation.group(1), equation.group(2), f"{equation_prefix}{equation_number}")
            continue
        if heading:
            flush()
            level = len(heading.group(1))
            title = _romanize_for_portable_pdf(heading.group(2))
            if level == 1:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _inline(paragraph, title)
            else:
                if title == "Figure legends":
                    document.add_page_break()
                if level == 2 and title in NUMBERED_SECTIONS:
                    current_section_number = NUMBERED_SECTIONS[title]
                    subsection_number = 0
                    title = f"{current_section_number}  {title}"
                elif level == 3 and current_section_number is not None:
                    subsection_number += 1
                    title = f"{current_section_number}.{subsection_number}  {title}"
                paragraph = document.add_paragraph(style=f"Heading {level - 1}")
                _inline(paragraph, title)
                in_references = title == "References"
            continue
        if re.match(r"^- ", line):
            flush()
            paragraph = document.add_paragraph(style="List Bullet")
            _inline(paragraph, _romanize_for_portable_pdf(line[2:]))
            continue
        if re.match(r"^\|", line):
            flush()
            continue
        if not line.strip():
            flush()
        else:
            if in_references and re.match(r"^\d+\.\s", line) and buffer:
                flush()
            buffer.append(line)
    flush()


def _add_figures(document: Document) -> None:
    document.add_page_break()
    heading = document.add_paragraph(style="Heading 1")
    heading.add_run("Figures")
    for number in range(1, 7):
        if number > 1:
            document.add_page_break()
        path = next((RESEARCH_ROOT / "figures").glob(f"fig{number}_*.png"))
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=Inches(6.55))
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = caption.add_run(f"Figure {number}")
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9)


def _add_tables(document: Document) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    heading = document.add_paragraph(style="Heading 1")
    heading.add_run("Supplementary Tables")
    for number in range(1, 13):
        if number > 1:
            document.add_page_break()
        title = document.add_paragraph(style="Heading 2")
        title.add_run(f"Supplementary Table S{number}")
        path = next((RESEARCH_ROOT / "manuscript/tables").glob(f"table{number}_*.csv"))
        with path.open(encoding="utf-8", newline="") as handle:
            rows = list(csv.reader(handle))
        if number == 8:
            label_column = rows[0].index("label_zh")
            rows = [[value for index, value in enumerate(row) if index != label_column] for row in rows]
        else:
            rows = [[_romanize_for_portable_pdf(value) for value in row] for row in rows]
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_borders(table)
        table.autofit = True
        for row_index, row in enumerate(rows):
            for col_index, value in enumerate(row):
                cell = table.cell(row_index, col_index)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                _set_cell_margins(cell)
                cell.text = value
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(6.5 if len(rows[0]) > 8 else 7.5)
                        run.font.color.rgb = INK
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
                if row_index == 0:
                    _set_header_bottom_border(cell)
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
                        run.font.color.rgb = INK
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def export() -> list[Path]:
    SUBMISSION_DIR.mkdir(exist_ok=True)
    build_complete_manuscript()

    manuscript = Document()
    _configure(manuscript, line_numbers=False)
    _add_markdown(manuscript, MAIN_MD)
    main_path = SUBMISSION_DIR / "xiangji_cgir_complete_paper.docx"
    manuscript.save(main_path)

    supplement = Document()
    _configure(supplement, line_numbers=False)
    _add_markdown(supplement, SUPPLEMENT_MD, stop_before_tables=True, equation_prefix="S")
    _add_tables(supplement)
    supplement_path = SUBMISSION_DIR / "xiangji_cgir_supplementary.docx"
    supplement.save(supplement_path)
    return [main_path, supplement_path]


def main() -> None:
    for path in export():
        print(path)


if __name__ == "__main__":
    main()
